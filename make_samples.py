"""Sample messages for testing the message library (invented wording, not real instructions)."""
import docx, pathlib
d = pathlib.Path('samples'); d.mkdir(exist_ok=True)
(d/'Aadhaar reminder - exact.txt').write_text(
"विषय: अभिभावक आधार सत्यापन के संबंध में।\n\nसमस्त जिला बेसिक शिक्षा अधिकारी,\n\nकृपया संलग्न रिपोर्ट के अनुसार लंबित आधार सत्यापन तीन दिवस में पूर्ण कराना सुनिश्चित करें।\n\nराज्य परियोजना निदेशक", encoding='utf-8')
doc = docx.Document()
for p in ["Subject: Progress on {indicator} – {period}", "All District Basic Education Officers,",
          "The state figure is {state_value}. {best_district} leads with {best_value}; {worst_district} is last with {worst_value}.",
          "Districts needing immediate attention:\n{bottom_list}", "Please act within one week."]:
    doc.add_paragraph(p)
doc.save(d/'Monthly review - figures.docx')
doc = docx.Document()
for p in ["समस्त जिला बेसिक शिक्षा अधिकारी,", "इस माह की स्थिति निम्नानुसार है।", "सभी जनपद अपनी प्रगति सुधारें तथा अनुपालन आख्या प्रेषित करें।"]:
    doc.add_paragraph(p)
doc.save(d/'Plain opening and closing.docx')
print(sorted(p.name for p in d.iterdir()))
