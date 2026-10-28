"""Test fixtures: a list of 75 placeholder district names, and a 'next period' file
whose heading for Total Enrollment is spelled differently (to test the mapping prompt)."""
import openpyxl, shutil
src='/mnt/user-data/uploads/Pm_SHRI_DBT_Verification.xlsx'
names=[f"परीक्षण जनपद {i:02d}" if i%2 else f"Test District {i:02d}" for i in range(1,76)]
open('test_district_list.txt','w',encoding='utf-8').write('\n'.join(names))
shutil.copy(src,'next_period.xlsx')
wb=openpyxl.load_workbook('next_period.xlsx'); ws=wb.active
ws['G2']='Total Enrolment'           # heading changed in the new export
for r in range(5,80): ws.cell(r,13).value = ws.cell(r,13).value + 10   # Guardian Verify moves
ws.cell(4,13).value = sum(ws.cell(r,13).value for r in range(5,80))
wb.save('next_period.xlsx')
